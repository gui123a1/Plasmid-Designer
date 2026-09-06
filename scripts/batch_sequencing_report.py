"""批量测序结果自动整理与分析（离线批处理，复用平台 Sanger 管线）

用法：
  python scripts/batch_sequencing_report.py --data-dir "<测序结果文件夹>" \
      [--excel "<信息表.xlsx>"] [--out-dir "<输出目录>"] [--min-q 20] [--dry-run]

流程：
  1. 读取 Excel 信息表（表头含"质粒名称/测序引物/测序结果"，引物列填测序文件名，
     多个用分号分隔，与平台网页端"参考文件与测序文件同目录导入"的约定一致）；
  2. 递归扫描数据文件夹中的 .ab1（测序 read）与 .dna/.gb/.fasta（参考图谱），
     按文件名主干与 Excel 引物列 / 质粒名称匹配，未匹配文件保留原位并列出清单；
  3. 把同一质粒的文件【复制】到 <out-dir>/<质粒名>/，原始文件一律不移动、
     不删除、不覆盖（整理清单记录原始路径与 MD5，保证可追溯）；
  4. 以同名图谱文件为参考，逐质粒运行 Sanger 全自动分析管线（峰级证据、
     置信度分层、CDS 结论），在质粒文件夹内生成 测序分析报告.md 与 分析结果.json；
  5. 结论回填 Excel 的"测序结果"列（原表先备份到输出目录）。
"""

import argparse
import csv
import hashlib
import json
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]
BACKEND = REPO_ROOT / "src" / "backend"
if (BACKEND / "core").is_dir():
    sys.path.insert(0, str(BACKEND))

READ_EXTS = {".ab1", ".ab"}
REF_EXTS = {".dna", ".gb", ".gbk", ".genbank", ".fasta", ".fa", ".fna"}
PRIMER_SPLIT = re.compile(r"[;；,，\n]+")


def _norm(s) -> str:
    """文件名/条目归一化：去空白、小写、去 .ab1 后缀，用于匹配。"""
    t = re.sub(r"\s+", "", str(s)).lower()
    for suf in (".ab1", ".ab"):
        if t.endswith(suf):
            t = t[: -len(suf)]
    return t.strip('"')


# ---------------------------------------------------------------- Excel 信息表


def load_excel(path: Path):
    """解析信息表，返回 (workbook, header_row, {列名: 列号}, [数据行])"""
    import openpyxl

    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    header_row, cols = None, {}
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20)):
        for c in row:
            if c.value and "质粒" in str(c.value):
                header_row = c.row
                break
        if header_row:
            break
    if header_row is None:
        raise SystemExit(f"Excel 中找不到含'质粒'的表头行：{path}")
    for c in ws[header_row]:
        v = str(c.value or "")
        if "质粒" in v:
            cols["plasmid"] = c.column
        elif "引物" in v:
            cols["primer"] = c.column
        elif "结果" in v:
            cols["result"] = c.column
    cols.setdefault("primer", cols["plasmid"] + 1)
    cols.setdefault("result", cols["primer"] + 1)

    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        name = ws.cell(row=r, column=cols["plasmid"]).value
        prim = ws.cell(row=r, column=cols["primer"]).value
        if name is None and prim is None:
            continue
        entries = [p.strip() for p in PRIMER_SPLIT.split(str(prim or "")) if p.strip()]
        rows.append({
            "row": r,
            "name": str(name or "").strip(),
            "primers": entries,
            "existing_result": ws.cell(row=r, column=cols["result"]).value,
        })
    return wb, header_row, cols, rows


# ---------------------------------------------------------------- 文件扫描与匹配


def scan_files(data_dir: Path, out_dir: Path):
    files = []
    for p in sorted(data_dir.rglob("*")):
        if not p.is_file() or p.name.startswith("~$"):
            continue
        if out_dir == p or out_dir in p.parents:
            continue
        ext = p.suffix.lower()
        if ext in READ_EXTS or ext in REF_EXTS:
            files.append({"path": p, "rel": p.relative_to(data_dir), "ext": ext,
                          "stem": _norm(p.stem)})
    return files


def match_files(rows, files):
    """按 Excel 引物列（=测序文件名）与质粒名称（=图谱文件名）匹配。

    返回 (per_plasmid, unmatched)。per_plasmid: 质粒名 -> {"reference":…, "reads":[…]}
    每个文件项为 {"file": 扫描项, "how": 匹配方式说明}
    """
    primer_map = {}  # 归一化文件名主干 -> (质粒名, 匹配方式)
    for row in rows:
        for e in row["primers"]:
            primer_map.setdefault(_norm(e), (row["name"], "Excel 引物列精确匹配"))

    per_plasmid: dict = {row["name"]: {"reference": None, "reads": []} for row in rows}
    unmatched = []

    for f in files:
        if f["ext"] in READ_EXTS:
            hit = primer_map.get(f["stem"])
            if hit is None:
                cands = {(p, how) for stem, (p, how) in primer_map.items() if stem in f["stem"]}
                if len(cands) == 1:
                    hit = next(iter(cands))
                    hit = (hit[0], f"文件名包含引物条目（{f['stem']}）")
            if hit is None:
                unmatched.append({"file": f, "reason": "未出现在任何质粒的引物列中"})
                continue
            per_plasmid.setdefault(hit[0], {"reference": None, "reads": []})["reads"].append(
                {"file": f, "how": hit[1]}
            )
        else:  # 图谱文件 → 按质粒名称匹配
            names = {_norm(r["name"]): r["name"] for r in rows if r["name"]}
            if f["stem"] in names:
                plasmid = names[f["stem"]]
                how = "文件名与质粒名称精确匹配"
            else:
                cands = {n for n in names if n and n in f["stem"]}
                if len(cands) != 1:
                    unmatched.append({"file": f, "reason": "文件名不匹配任何质粒名称"})
                    continue
                plasmid = names[cands.pop()]
                how = "文件名包含质粒名称"
            slot = per_plasmid.setdefault(plasmid, {"reference": None, "reads": []})
            if slot["reference"] is not None:
                unmatched.append({"file": f,
                                  "reason": f"质粒 {plasmid} 已有图谱 {slot['reference']['file']['path'].name}，本文件未采用"})
                continue
            slot["reference"] = {"file": f, "how": how}
    return per_plasmid, unmatched


# ---------------------------------------------------------------- 文件整理（只复制、不覆盖）


def md5_of(path: Path) -> str:
    h = hashlib.md5()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def organize(plasmid: str, items, out_dir: Path, manifest, dry_run=False):
    """把属于该质粒的文件复制到 <out-dir>/<质粒名>/，返回报告用的清单。"""
    folder = out_dir / plasmid
    copied = []
    for it in items:
        src = it["file"]["path"]
        dst = folder / src.name
        n = 2
        while dst.exists():  # 永不覆盖：同名加序号
            dst = folder / f"{src.stem} ({n}){src.suffix}"
            n += 1
        if not dry_run:
            folder.mkdir(parents=True, exist_ok=True)
            shutil.copy2(src, dst)
        entry = {
            "plasmid": plasmid, "name": src.name,
            "kind": "参考图谱" if it["file"]["ext"] in REF_EXTS else "测序 read",
            "how": it["how"], "src": str(src), "dst": str(dst),
            "kb": round(src.stat().st_size / 1024, 1), "md5": md5_of(src),
        }
        manifest.append(entry)
        copied.append(entry)
    return copied


# ---------------------------------------------------------------- 分析与报告


def analyze_plasmid(reference: Path, reads, min_q: int):
    from core.sanger.pipeline import analyze
    from core.sanger.reference_parser import parse_reference

    ref_seq, features = parse_reference(reference.name, reference.read_bytes())
    ab1s = [(p["file"]["path"].name, p["file"]["path"].read_bytes()) for p in reads]
    res = analyze(ab1s, ref_seq, features, min_q=min_q)
    res["_ref_len"] = len(ref_seq)
    res["_n_features"] = len(features)
    return res


def _fmt_change(v) -> str:
    t = v["type"]
    if t == "substitution":
        return f"{v.get('ref_base', '?')}>{v.get('alt_base', '?')}"
    if t == "deletion":
        return f"Δ{v.get('length', 1)}bp"
    if t == "insertion":
        return f"+{v.get('alt_base', '?')}"
    return t


def _fmt_evidence(v) -> str:
    ev = v.get("peak_evidence") or {}
    parts = []
    if ev.get("mutant_pct") is not None:
        parts.append(f"峰强比 {ev['mutant_pct']:.0f}%")
    if ev.get("snr") is not None:
        parts.append(f"信噪比 {ev['snr']:.1f}")
    if ev.get("insertion_peak_ratio") is not None:
        parts.append(f"插入峰强度 {ev['insertion_peak_ratio'] * 100:.0f}%")
    if v.get("corroborated_by_basecall"):
        parts.append("重basecall印证")
    return "、".join(parts) or "—"


def main_sentence(verdict: str) -> str:
    """从 CDS verdict 中取核心主句（去掉覆盖度前缀与待复核尾句）。"""
    seg = verdict.split("。")
    return seg[1] if len(seg) > 1 else verdict


def excel_conclusion(plasmid: str, res, n_reads: int, has_ref: bool) -> str:
    """回填到 Excel'测序结果'列的简要结论（一句话，可直接行动）。"""
    if not has_ref:
        return f"已整理 {n_reads} 个测序文件，但缺少同名参考图谱（.dna/.gb/.fasta），无法自动比对"
    if res is None:
        return "分析失败"
    if not res["reads"]:
        errs = "；".join(e.get("error", "") for e in res["errors"][:2]) or "无可分析 read"
        return f"分析失败：{errs}"
    confirmed = [v for v in res["variants"] if v.get("confidence") != "low"]
    pending = len(res["variants"]) - len(confirmed)
    cov = res["consensus"]["coverage_percent"]
    bad = [c for c in res["cds_reports"]
           if c["coverage_status"] != "uncovered" and c["protein_identical"] is False]
    tail = []
    if pending:
        tail.append(f"{pending} 处低置信差异疑似测序噪声（详见报告）")
    if cov < 95:
        tail.append(f"测序覆盖 {cov:.0f}%")

    if bad:
        body = "；".join(f"[{c['name']}] {main_sentence(c['verdict'])}" for c in bad[:2])
        if len(bad) > 2:
            body += f"（另有 {len(bad) - 2} 个 CDS 不一致）"
        out = f"不合格：{body}"
    elif confirmed:
        feats = "、".join(sorted({f["name"] for v in confirmed for f in v.get("features", [])})) or "无特征区"
        out = (f"合格：编码区蛋白与设计一致；编码区外有 {len(confirmed)} 处确证差异"
               f"（位于 {feats}），不影响编码产物")
    else:
        out = "合格：与设计一致"
    if tail:
        out += "；" + "；".join(tail)
    return out


def write_report_md(plasmid: str, out_path: Path, ctx: dict, res, copied):
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    engine_note = ""
    if res and res["engine"] == "internal+biopython":
        engine_note = "；本机未安装 tracy，跳过重 basecall 交叉印证"
    lines = [f"# {plasmid} 测序分析报告", "", f"- 生成时间：{now}",
             f"- 分析引擎：{res['engine']}（平台 Sanger 全自动管线，本地批处理{engine_note}）" if res else "- 未运行序列分析",
             f"- 参考图谱：{ctx['reference_name']}（{res['_ref_len']} bp，{res['_n_features']} 个特征）" if res
             else f"- 参考图谱：{ctx['reference_name'] or '缺失'}",
             ""]
    lines += ["## 结论", ""]
    concl = excel_conclusion(plasmid, res, len(ctx["reads"]), ctx["reference"] is not None)
    lines.append(f"> **{concl}**")
    lines.append("")
    if res and res["reads"]:
        lines += ["各编码区（CDS）结论：", ""]
        for c in res["cds_reports"]:
            mark = {"full": "✅", "partial": "🟡", "uncovered": "⚪"}.get(c["coverage_status"], "")
            lines.append(f"- {mark} **{c['name']}**（{c['covered_percent']}% 覆盖）：{c['verdict']}")
        lines.append("")
    lines += ["## 文件整理（原始文件未改动，此处为副本）", "",
              "| 文件 | 类型 | 匹配方式 | 原始路径 |", "|---|---|---|---|"]
    for e in copied:
        lines.append(f"| {e['name']} | {e['kind']} | {e['how']} | {e['src']} |")
    lines.append("")
    if res and res["reads"]:
        lines += ["## 测序 read 概况", "",
                  "| 文件 | QC 等级 | 平均 Q | 有效长度 | 比对区段（参考坐标） |",
                  "|---|---|---|---|---|"]
        for r in res["reads"]:
            a = r["alignment"]
            lines.append(f"| {r['filename']} | {r['grade']} | {r['mean_q']} | {r['trimmed_length']} "
                         f"| {a['ref_start']}–{a['ref_end']} |")
        lines.append("")
    if res and res["variants"]:
        lines += ["## 变异明细", "",
                  "| 位置 | 类型 | 变化 | 置信度 | 测序Q | 支持read数 | 峰级证据 | 所在特征 | 氨基酸变化 |",
                  "|---|---|---|---|---|---|---|---|---|"]
        for v in res["variants"]:
            feats = "、".join(f["name"] for f in v.get("features", [])) or "—"
            aa = v.get("aa_change") or ("同义" if v.get("synonymous")
                                        else "移码" if v.get("frameshift") else "—")
            conf = {"high": "高", "medium": "中", "low": "低"}.get(v.get("confidence"), v.get("confidence"))
            lines.append(f"| {v['ref_pos']} | {v['type']} | {_fmt_change(v)} | {conf} "
                         f"| {v.get('read_q', '')} | {v.get('support_reads', 1)} "
                         f"| {_fmt_evidence(v)} | {feats} | {aa} |")
        lines += ["", "置信度说明：**高**=峰级证据与 Q 值充分支持；**中**=证据一般，建议关注；"
                  "**低**=疑似测序噪声，未计入结论。峰强比=突变碱基峰占双峰百分比（≈100% 纯合、"
                  "≈50% 混合）；插入峰强度=插入碱基峰面积÷邻峰面积中位数（≥60% 说明插入峰真实存在）。",
                  ""]
    if res and res["errors"]:
        lines += ["## 分析中的错误", ""]
        lines += [f"- {e['filename']}: {e['error']}" for e in res["errors"]]
        lines.append("")
    out_path.write_text("\n".join(lines), encoding="utf-8")


# ---------------------------------------------------------------- 主流程


def main():
    ap = argparse.ArgumentParser(description="批量测序结果自动整理与分析")
    ap.add_argument("--data-dir", required=True, help="测序结果文件夹（含 .ab1/.dna 与信息表）")
    ap.add_argument("--excel", help="信息表路径（默认取文件夹下第一个 .xlsx）")
    ap.add_argument("--out-dir", help="输出目录（默认 <data-dir>/测序分析）")
    ap.add_argument("--min-q", type=int, default=20, help="修剪 Q 阈值（默认 20）")
    ap.add_argument("--dry-run", action="store_true", help="只匹配与预览，不复制、不分析、不回填")
    ap.add_argument("--no-write-back", action="store_true", help="不回填 Excel 结论")
    args = ap.parse_args()

    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    data_dir = Path(args.data_dir).resolve()
    if not data_dir.is_dir():
        raise SystemExit(f"数据文件夹不存在：{data_dir}")
    excel_path = Path(args.excel).resolve() if args.excel else next(
        (p for p in sorted(data_dir.glob("*.xlsx")) if not p.name.startswith("~$")), None)
    if excel_path is None:
        raise SystemExit(f"数据文件夹下找不到 Excel 信息表（.xlsx）：{data_dir}")
    out_dir = Path(args.out_dir).resolve() if args.out_dir else data_dir / "测序分析"

    wb, header_row, cols, rows = load_excel(excel_path)
    print(f"信息表：{excel_path.name}（{len(rows)} 个质粒）")
    files = scan_files(data_dir, out_dir)
    n_ab1 = sum(1 for f in files if f["ext"] in READ_EXTS)
    print(f"扫描到 {len(files)} 个文件（{n_ab1} 个 .ab1，{len(files) - n_ab1} 个图谱）")

    per_plasmid, unmatched = match_files(rows, files)

    if args.dry_run:
        for name, slot in per_plasmid.items():
            ref = slot["reference"]["file"]["path"].name if slot["reference"] else "（缺）"
            print(f"[预览] {name}: 图谱={ref}, reads={[r['file']['path'].name for r in slot['reads']]}")
        for u in unmatched:
            print(f"[未匹配] {u['file']['path'].name}: {u['reason']}")
        return

    manifest, conclusions = [], {}
    for row in rows:
        name = row["name"]
        slot = per_plasmid.get(name, {"reference": None, "reads": []})
        items = ([slot["reference"]] if slot["reference"] else []) + slot["reads"]
        copied = organize(name, items, out_dir, manifest) if items else []
        if not slot["reads"] and slot["reference"] is None:
            conclusions[name] = "未在文件夹中找到该质粒的测序文件或图谱"
            print(f"· {name}: 未找到文件")
            continue

        res = None
        if slot["reference"] and slot["reads"]:
            try:
                res = analyze_plasmid(slot["reference"]["file"]["path"],
                                      slot["reads"], args.min_q)
            except Exception as e:  # noqa: BLE001 单个质粒失败不阻断批次
                conclusions[name] = f"分析失败：{e}"
                print(f"· {name}: 分析失败 {e}")
                continue
        concl = excel_conclusion(name, res, len(slot["reads"]), slot["reference"] is not None)
        conclusions[name] = concl

        folder = out_dir / name
        if res is not None:
            report = {k: v for k, v in res.items() if k != "traces"}
            report["_generated_at"] = datetime.now().isoformat(timespec="seconds")
            report["_reference_file"] = slot["reference"]["file"]["path"].name
            (folder / "分析结果.json").write_text(
                json.dumps(report, ensure_ascii=False, indent=1), encoding="utf-8")
        write_report_md(name, folder / "测序分析报告.md",
                        {"reference": slot["reference"], "reference_name":
                         slot["reference"]["file"]["path"].name if slot["reference"] else None,
                         "reads": slot["reads"]}, res, copied)
        n_var = len(res["variants"]) if res else 0
        print(f"· {name}: {len(slot['reads'])} read、{n_var} 处变异 → {concl[:60]}…")

    for u in unmatched:
        print(f"! 未匹配（保留原位）：{u['file']['path'].name} — {u['reason']}")

    # 整理清单（可追溯：原始路径 + MD5）
    if manifest:
        with open(out_dir / "整理清单.csv", "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.DictWriter(fh, fieldnames=list(manifest[0].keys()))
            w.writeheader()
            w.writerows(manifest)

    # Excel 备份 + 回填
    if args.no_write_back:
        print("已跳过 Excel 回填（--no-write-back）")
        return
    backup = out_dir / f"{excel_path.stem}_原始备份{excel_path.suffix}"
    n = 2
    while backup.exists():
        backup = out_dir / f"{excel_path.stem}_原始备份({n}){excel_path.suffix}"
        n += 1
    out_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy2(excel_path, backup)
    for row in rows:
        concl = conclusions.get(row["name"])
        if not concl:
            continue
        cell = wb.worksheets[0].cell(row=row["row"], column=cols["result"])
        old = str(cell.value or "").strip()
        cell.value = concl if not old or old == concl else f"{old}\n{concl}"
    wb.save(excel_path)
    print(f"√ 结论已回填 {excel_path.name}（原表备份：{backup.name}）")
    print(f"√ 输出目录：{out_dir}")


if __name__ == "__main__":
    main()
