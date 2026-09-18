"""批量测序结果归组与一句话结论（网页端批量分析与离线脚本共用）

测序公司交付常为「Excel 信息表 + 一堆 .ab1/.dna」：信息表的引物列填测序
文件名（可多个，中英文分号/逗号均可分隔），质粒名称列对应参考图谱文件名主干。

两种信息表形态：
- 质粒模式（match_files）：无克隆号列，引物列填测序文件名主干；
- 克隆模式（match_clone_files）：带「克隆号」列，一个质粒多个克隆每行一个，
  引物列填引物名，ab1 文件名 = 「克隆号-引物」（如 S99678-M13F-75.ab1），
  每个克隆独立成组、独立分析；质粒名称与图谱文件名支持两段式部分匹配
  （"123-1 AB2C" 可只写 "123-1"、"AB2C" 或全名，候选唯一才采纳）。

本模块只做与入口无关的纯逻辑：
- 解析 Excel 信息表（路径或字节流均可）；
- 按引物列/质粒名称/克隆号把扫描到的文件归组；
- 从单样品分析结果生成一句话结论。

文件整理（复制归档、MD5 清单）、报告输出与 Excel 回填属于各入口自身
的流程：离线批处理见 scripts/batch_sequencing_report.py，网页批量分析见
POST /api/sequencing/analyze-batch。
"""

import io
import re
from typing import Dict, List, Optional, Tuple, Union

READ_EXTS = {".ab1", ".ab"}
REF_EXTS = {".dna", ".gb", ".gbk", ".genbank", ".fasta", ".fa", ".fna"}
PRIMER_SPLIT = re.compile(r"[;；,，\n]+")


def norm_stem(name) -> str:
    """文件名/条目归一化：去空白、小写、去 .ab1 后缀，用于匹配。"""
    t = re.sub(r"\s+", "", str(name)).lower()
    for suf in (".ab1", ".ab"):
        if t.endswith(suf):
            t = t[: -len(suf)]
    return t.strip('"')


def _squash(name) -> str:
    """分隔符不敏感形式：小写并去掉空白与 -/_。'S99678-M13F-75' → 's99678m13f75'"""
    return re.sub(r"[\s\-_]+", "", str(name).lower())


def _cell_str(v) -> str:
    """Excel 单元格 → 字符串；数值 123.0 → '123'（克隆号/编号不能带 .0）"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def load_excel(src: Union[str, "Path", bytes, bytearray]):
    """解析信息表，返回 (workbook, header_row, {列名: 列号}, [数据行])

    src：文件路径（脚本离线用）或字节流（网页上传用）。
    表头行取第一个含"质粒"的单元格所在行；找不到抛 ValueError。
    """
    import openpyxl

    if isinstance(src, (bytes, bytearray)):
        wb = openpyxl.load_workbook(io.BytesIO(src))
    else:
        wb = openpyxl.load_workbook(src)
    ws = wb.worksheets[0]
    header_row, cols = None, {}
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20)):
        for c in row:
            if c.value and "质粒" in str(c.value):
                header_row = c.row
                break
        if header_row:
            break
    if header_row is None:
        raise ValueError("Excel 中找不到含'质粒'的表头行")
    for c in ws[header_row]:
        v = str(c.value or "")
        if "质粒" in v:
            cols["plasmid"] = c.column
        elif "引物" in v:
            cols["primer"] = c.column
        elif "结果" in v:
            cols["result"] = c.column
        elif "克隆" in v or "样品" in v:
            cols["clone"] = c.column
    cols.setdefault("primer", cols["plasmid"] + 1)
    cols.setdefault("result", cols["primer"] + 1)

    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        name = _cell_str(ws.cell(row=r, column=cols["plasmid"]).value)
        prim = _cell_str(ws.cell(row=r, column=cols["primer"]).value)
        if not name and not prim:
            continue
        entries = [p.strip() for p in PRIMER_SPLIT.split(prim) if p.strip()]
        rows.append({
            "row": r,
            "name": name,
            "clone": _cell_str(ws.cell(row=r, column=cols["clone"]).value) if "clone" in cols else "",
            "primers": entries,
            "existing_result": ws.cell(row=r, column=cols["result"]).value,
        })
    return wb, header_row, cols, rows


def match_files(rows: List[Dict], files: List[Dict]):
    """按 Excel 引物列（=测序文件名）与质粒名称（=图谱文件名）匹配。

    files 的每项为 {"ext": ".ab1"/".dna"/…, "stem": 归一化主干, "name"(可选):
    显示文件名, "path"(可选): 离线脚本场景的完整路径}——至少有其一。

    返回 (per_plasmid, unmatched)。per_plasmid: 质粒名 -> {"reference":…, "reads":[…]}
    每个文件项为 {"file": 原文件项, "how": 匹配方式说明}
    """

    def fname(f: Dict) -> str:
        return f.get("name") or str(f["path"].name)

    primer_map = {}  # 归一化文件名主干 -> (质粒名, 匹配方式)
    for row in rows:
        for e in row["primers"]:
            primer_map.setdefault(norm_stem(e), (row["name"], "Excel 引物列精确匹配"))

    per_plasmid: Dict[str, Dict] = {row["name"]: {"reference": None, "reads": []} for row in rows}
    unmatched: List[Dict] = []

    for f in files:
        if f["ext"] in READ_EXTS:
            hit = primer_map.get(f["stem"])
            if hit is None:
                cands = {(p, how) for stem, (p, how) in primer_map.items() if stem in f["stem"]}
                if len(cands) == 1:
                    hit = next(iter(cands))
                    hit = (hit[0], f"文件名包含引物条目（{f['stem']}）")
            if hit is None:
                unmatched.append({"file": f, "reason": "未出现在任何质粒的引物列中"})
                continue
            per_plasmid.setdefault(hit[0], {"reference": None, "reads": []})["reads"].append(
                {"file": f, "how": hit[1]}
            )
        else:  # 图谱文件 → 按质粒名称匹配
            names = {norm_stem(r["name"]): r["name"] for r in rows if r["name"]}
            if f["stem"] in names:
                plasmid = names[f["stem"]]
                how = "文件名与质粒名称精确匹配"
            else:
                cands = {n for n in names if n and n in f["stem"]}
                if len(cands) != 1:
                    unmatched.append({"file": f, "reason": "文件名不匹配任何质粒名称"})
                    continue
                plasmid = names[cands.pop()]
                how = "文件名包含质粒名称"
            slot = per_plasmid.setdefault(plasmid, {"reference": None, "reads": []})
            if slot["reference"] is not None:
                unmatched.append({"file": f,
                                  "reason": f"质粒 {plasmid} 已有图谱 {fname(slot['reference']['file'])}，本文件未采用"})
                continue
            slot["reference"] = {"file": f, "how": how}
    return per_plasmid, unmatched


def rows_have_clones(rows: List[Dict]) -> bool:
    """信息表带「克隆号」列且至少一行有值 → 克隆模式（每行一个克隆，独立分析）"""
    return any(r.get("clone") for r in rows)


def match_clone_files(rows: List[Dict], files: List[Dict]):
    """克隆号模式归组：一个质粒多个克隆，每行（克隆）独立成组、独立分析。

    files 的每项为 {"ext", "stem", "name"(可选), "path"(可选)}（同 match_files）。

    read 匹配优先级：
    1. 「克隆号-引物」组合精确（含忽略分隔符形式，S99678_M13F-75 也能命中）；
    2. 引物列条目本身精确（兼容引物列直接填完整文件名的旧表）；
    3. 主干包含某期望组合（复制版后缀 -2/_copy 等）；
    4. 主干以某克隆号开头（克隆号后是分隔符或结尾）——信息表漏填引物、
       或一个克隆只有一个文件时兜底。
    同一引物名出现在多个克隆行时，只写引物名的文件无法归属，判为未匹配。

    图谱 → 质粒名称支持两段式部分匹配：全名精确 > 忽略分隔符精确 >
    与名称的某一段一致（"MBYSTC" ↔ "17648 MBYSTC"）> 双向包含（短侧 ≥4 字符，
    防 "MX" 命中 "MX2"）；候选不唯一一律不猜，判为未匹配。同一质粒的不同
    克隆行可混用全名与单段名（'17648' / 'MBYSTC' / '17648 MBYSTC'）——只写
    一段的名称按双向包含找唯一候选图谱共享（候选多于一张维持缺图谱）。
    图谱文件名也可只写一段（'17648.fasta'）：与图谱共享该段的行直接命中；
    只写另一段的行（'MBYSTC'）与之没有共同标识，仅当表内恰有一个以它为
    组成段的质粒名且该名已确定图谱时，按表内别名共享；表中无全名行时
    确实无法建立对应关系，维持缺图谱。

    返回 (groups, unmatched)。groups 顺序 = 信息表行序，每项：
    {"plasmid", "clone", "reference": {"file","how"}|None, "reads": [{"file","how"}]}
    """
    groups: Dict[Tuple[str, str], Dict] = {}
    for r in rows:
        key = (r["name"], r.get("clone", ""))
        if key not in groups:
            groups[key] = {"plasmid": r["name"], "clone": r.get("clone", ""),
                           "reference": None, "reads": []}
    names = {r["name"] for r in rows if r["name"]}

    # ---- 图谱 → 质粒名称（返回 质粒名, 说明, 匹配强度 0=全名精确 1=忽略分隔符
    #      精确 2=与名称某一段一致 3=双向包含）----
    def _match_ref(f: Dict) -> Tuple[Optional[str], str, int]:
        stem, sq = f["stem"], _squash(f["stem"])
        levels = (
            ([n for n in names if norm_stem(n) == stem],
             "文件名与质粒名称精确匹配"),
            ([n for n in names if _squash(n) == sq],
             "文件名与质粒名称精确匹配（忽略分隔符）"),
            ([n for n in names if sq in {_squash(t) for t in re.split(r"\s+", n.strip()) if t}],
             None),
        )
        cont = []
        for n in names:
            nsq = _squash(n)
            if min(len(nsq), len(sq)) >= 4 and (nsq in sq or sq in nsq):
                cont.append(n)
        for level, (cands, how) in enumerate(levels):
            if len(cands) == 1:
                return cands[0], (how or f"文件名与质粒名称的组成部分一致（{cands[0]}）"), level
            if len(cands) > 1:
                return None, ("文件名同时匹配多个质粒名称（"
                              + "、".join(sorted(cands))[:80] + "），无法唯一确定"), level
        if len(cont) == 1:
            return cont[0], "文件名包含质粒名称（或反之）", 3
        if len(cont) > 1:
            return None, ("文件名同时匹配多个质粒名称（"
                          + "、".join(sorted(cont))[:80] + "），无法唯一确定"), 3
        return None, "文件名不匹配任何质粒名称（支持全名或两段式名称的任一段）", 3

    # ---- 期望 read 文件名主干：克隆号-引物 组合 + 引物条目本身 ----
    composed: Dict[str, Tuple[str, str, str]] = {}   # 主干 -> (质粒, 克隆, 匹配方式)
    composed_conflicts = set()                       # 同组合对应多行（如克隆号重复于不同质粒）
    bare: Dict[str, List[Tuple[str, str]]] = {}      # 引物条目 -> [(质粒, 克隆)]
    for r in rows:
        plasmid, clone = r["name"], r.get("clone", "")
        for e in r["primers"]:
            ne = norm_stem(e)
            if not ne:
                continue
            if clone:
                ck = f"{norm_stem(clone)}-{ne}"
                if composed.get(ck, (plasmid, clone))[:2] != (plasmid, clone):
                    composed_conflicts.add(ck)
                composed.setdefault(ck, (plasmid, clone, "克隆号+引物组合精确匹配"))
            bare.setdefault(ne, []).append((plasmid, clone))
    bare_unique = {k: v[0] for k, v in bare.items() if len(v) == 1}
    bare_ambiguous = {k for k, v in bare.items() if len(v) > 1}

    sq_pool: Dict[str, List[Tuple[str, str, str]]] = {}
    for k, v in composed.items():
        if k not in composed_conflicts:
            sq_pool.setdefault(_squash(k), []).append(v)
    for k, (p, c) in bare_unique.items():
        sq_pool.setdefault(_squash(k), []).append((p, c, "引物列精确匹配"))

    def _assign_read(f: Dict) -> Tuple[Optional[Tuple[str, str]], str]:
        stem, sq = f["stem"], _squash(f["stem"])
        if stem in composed_conflicts:
            return None, "「克隆号+引物」组合在信息表中对应多行，无法唯一归属"
        hit = composed.get(stem)
        if hit:
            return (hit[0], hit[1]), hit[2]
        b = bare_unique.get(stem)
        if b:
            return b, "引物列精确匹配"
        hits = sq_pool.get(sq)
        if hits and len(hits) == 1:
            return (hits[0][0], hits[0][1]), hits[0][2]
        if hits:
            return None, "文件名同时匹配多个「克隆号+引物」组合，无法唯一归属"
        # 兜底 1：主干包含某「克隆号+引物」组合（复制版/带额外后缀）。
        # 裸引物条目只做精确匹配、不参与包含——含引物名但克隆号未知/不在表中
        # 的文件（如 S99999-M13F-75）不能凭引物名吸收进某个克隆
        cont = [(k, v) for k, v in composed.items()
                if k not in composed_conflicts and len(_squash(k)) >= 4 and _squash(k) in sq]
        if cont:
            uniq = {(v[0], v[1]) for _, v in cont}
            if len(uniq) == 1:
                best = max(cont, key=lambda kv: len(kv[0]))[1]
                return (best[0], best[1]), "文件名包含「克隆号+引物」组合"
            return None, "文件名同时匹配多个克隆的引物组合，无法唯一归属"
        if stem in bare_ambiguous:
            return None, "该引物名出现在多个克隆行，无法确定克隆"
        # 兜底 2：主干以某克隆号开头（分隔符或结尾边界），信息表漏填引物时兜底
        for r in rows:
            c = r.get("clone", "")
            nc = norm_stem(c)
            if len(nc) < 2 or not stem.startswith(nc):
                continue
            nxt = stem[len(nc):len(nc) + 1]
            if nxt and nxt not in "-_ .":
                continue
            return (r["name"], c), f"文件名以克隆号 {c} 开头"
        return None, "未在任何克隆行的引物列中，也匹配不到「克隆号+引物」组合"

    unmatched: List[Dict] = []
    ref_by_plasmid: Dict[str, Dict] = {}
    # 同一质粒名可能被多张图谱文件声明：按匹配强度取最精确者；同强度并列时
    # 仅"文件主干一致"（全名/忽略分隔符精确）视为同一图谱的重复副本取其一，
    # 其余（如两段式部分匹配命中不同文件）无法确定采用哪张，全部不猜
    ref_claims: Dict[str, List[Tuple[Dict, str, int]]] = {}
    for f in files:
        if f["ext"] in READ_EXTS:
            key, how = _assign_read(f)
            if key is None:
                unmatched.append({"file": f, "reason": how})
            else:
                groups.setdefault(key, {"plasmid": key[0], "clone": key[1],
                                        "reference": None, "reads": []})
                groups[key]["reads"].append({"file": f, "how": how})
        else:
            plasmid, how, level = _match_ref(f)
            if plasmid is None:
                unmatched.append({"file": f, "reason": how})
            else:
                ref_claims.setdefault(plasmid, []).append((f, how, level))
    for plasmid, claims in ref_claims.items():
        claims.sort(key=lambda c: c[2])
        best_f, best_how, best_level = claims[0]
        tied = [c for c in claims if c[2] == best_level]
        if len(tied) > 1 and best_level >= 2:
            names_list = "、".join(c[0].get("name") or str(c[0]["path"].name) for c in claims)
            for c in claims:
                unmatched.append({"file": c[0],
                                  "reason": f"同时匹配质粒 {plasmid} 的多张候选图谱（{names_list[:80]}），无法唯一确定"})
            continue
        ref_by_plasmid[plasmid] = {"file": best_f, "how": best_how}
        for c in claims[1:]:
            other = c[0].get("name") or str(c[0]["path"].name)
            unmatched.append({"file": c[0],
                              "reason": f"质粒 {plasmid} 已采用图谱 {best_f.get('name') or str(best_f['path'].name)}，本文件未采用"})

    # 名称驱动兜底：表里同一质粒的不同克隆可能分别只写两段式名称的一段
    # （'17648'/'MBYSTC' 与全名 '17648 MBYSTC' 混用）——只写一段的质粒名
    # 在所有图谱文件中找唯一的双向包含候选，命中则共享该图谱；候选多于
    # 一张（如 'AB2C' 同时命中 '123-1 AB2C' 与 '123-2 AB2C'）维持缺图谱不猜
    ref_files = [f for f in files if f["ext"] not in READ_EXTS]
    for (plasmid, _clone) in list(groups):
        if not plasmid or plasmid in ref_by_plasmid:
            continue
        nsq = _squash(plasmid)
        if len(nsq) < 4:
            continue
        cands = [f for f in ref_files
                 if min(len(nsq), len(_squash(f["stem"]))) >= 4
                 and (nsq in _squash(f["stem"]) or _squash(f["stem"]) in nsq)]
        if len(cands) == 1:
            ref_by_plasmid[plasmid] = {
                "file": cands[0],
                "how": "质粒名称与图谱文件名部分一致（唯一候选，共享全名质粒的图谱）",
            }

    # 表内别名兜底：图谱文件名只含两段式名称的一段时（如图谱叫 '17648'），
    # 只写另一段的行（'MBYSTC'）与图谱没有共同标识——但信息表内若恰有一个
    # 以它为组成段的质粒名且该名已确定图谱，说明两种写法指向同一质粒，共享
    # 其图谱；这样的宿主名称多于一个（'MBYSTC' 同时是 17648/17649 的组成段）
    # 或宿主本身缺图谱时维持不猜。表中完全没有全名行时（图谱 '17648' ↔
    # 表 'MBYSTC'）确实无法建立对应，保持缺图谱
    for (plasmid, _clone) in list(groups):
        if not plasmid or plasmid in ref_by_plasmid:
            continue
        nsq = _squash(plasmid)
        if len(nsq) < 2:
            continue
        hosts = [n for n in names if n != plasmid and nsq in
                 {_squash(t) for t in re.split(r"\s+", n.strip()) if t}]
        if len(hosts) != 1:
            continue
        host_ref = ref_by_plasmid.get(hosts[0])
        if host_ref:
            ref_by_plasmid[plasmid] = {
                "file": host_ref["file"],
                "how": f"「{plasmid}」与「{hosts[0]}」为同一质粒的不同写法（表内互为别名），共享其图谱",
            }

    ordered = list(groups.values())
    for g in ordered:
        g["reference"] = ref_by_plasmid.get(g["plasmid"])
    return ordered, unmatched


def main_sentence(verdict: str) -> str:
    """从 CDS verdict 中取核心主句（去掉覆盖度前缀与待复核尾句）。"""
    seg = verdict.split("。")
    return seg[1] if len(seg) > 1 else verdict


def excel_conclusion(plasmid: str, res, n_reads: int, has_ref: bool) -> str:
    """一句话结论（网页结果表与离线 Excel 回填共用同一口径）。"""
    if not has_ref:
        return f"已整理 {n_reads} 个测序文件，但缺少同名参考图谱（.dna/.gb/.fasta），无法自动比对"
    if res is None:
        return "分析失败"
    if not res["reads"]:
        errs = "；".join(e.get("error", "") for e in res["errors"][:2]) or "无可分析 read"
        return f"分析失败：{errs}"
    confirmed = [v for v in res["variants"] if v.get("confidence") != "low"]
    pending = len(res["variants"]) - len(confirmed)
    cov = res["consensus"]["coverage_percent"]
    bad = [c for c in res["cds_reports"]
           if c["coverage_status"] != "uncovered" and c["protein_identical"] is False]
    tail = []
    if pending:
        tail.append(f"{pending} 处低置信差异疑似测序噪声（详见报告）")
    # A4：poly 区峰图计数与碱基调用不一致——此前批量 Excel 丢掉这条关键告警
    hp_bad = [e for e in (res.get("homopolymers") or [])
              if e.get("tier", "poly") == "poly" and e.get("count_reliable") is False]
    if hp_bad:
        tail.append("poly 区峰图计数与碱基调用不一致（重复数存疑，详见报告）")
    if cov < 95:
        tail.append(f"测序覆盖 {cov:.0f}%")

    if bad:
        body = "；".join(f"[{c['name']}] {main_sentence(c['verdict'])}" for c in bad[:2])
        if len(bad) > 2:
            body += f"（另有 {len(bad) - 2} 个 CDS 不一致）"
        out = f"不合格：{body}"
    elif confirmed:
        feats = "、".join(sorted({f["name"] for v in confirmed for f in v.get("features", [])})) or "无特征区"
        out = (f"合格：编码区蛋白与设计一致；编码区外有 {len(confirmed)} 处确证差异"
               f"（位于 {feats}），不影响编码产物")
    else:
        out = "合格：与设计一致"
    if tail:
        out += "；" + "；".join(tail)
    return out
